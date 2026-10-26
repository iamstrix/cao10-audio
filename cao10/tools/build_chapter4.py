"""Append measured Chapter 4 to the original proposal, preserving its contents."""
import json
from pathlib import Path
import re
import shutil
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches,Pt,RGBColor
ROOT=Path(__file__).resolve().parents[1]
WORK=ROOT.parent
target=next(WORK.glob('*.docx'))
backup=WORK/'tmp/docx-original/proposal-before-chapter4.docx'
backup.parent.mkdir(parents=True,exist_ok=True)
if not backup.exists():shutil.copy2(target,backup)
doc=Document(backup)
# Original table paragraphs inherit negative indents, clipping wrapped lines.
# Correct their cell layout without changing any of the proposal's text.
for existing in doc.tables:
    for row in existing.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.left_indent=Inches(0)
                paragraph.paragraph_format.right_indent=Inches(0)
                paragraph.paragraph_format.first_line_indent=Inches(0)
for paragraph in doc.paragraphs:
    pr=paragraph._p.pPr
    if pr is not None and pr.find(qn('w:shd')) is not None and any(r.font.name=='Roboto Mono' for r in paragraph.runs):
        paragraph.paragraph_format.left_indent=Inches(0)
        paragraph.paragraph_format.right_indent=Inches(0)
        paragraph.paragraph_format.first_line_indent=Inches(0)
        for run in paragraph.runs:
            run.font.name='Consolas'
            run.font.size=Pt(9)
def read(name):return json.loads((ROOT/'results'/name).read_text(encoding='utf-8-sig'))
native=read('engine-tests.json');headless=read('headless.json');gui=read('gui-headless.json')
bench=read('benchmark.json');env=read('environment.json');cpu=read('cpu.json')
waves=['Square','Sawtooth','Triangle','Noise']
lookup={(r['implementation'],r['waveform']):r for r in bench}
# Count statically listed instructions across all alternative paths in the loop.
listing=(ROOT/'build/sfx.lst').read_text()
region=listing.split('.loop:')[1].split('.done:')[0]
loop_instructions=sum(bool(re.match(r'^\s*\d+\s+[0-9A-F]{8}\s+[0-9A-F]',line)) for line in region.splitlines())

def style(name,size=12,bold=False):
    s=doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    s.base_style=doc.styles['normal'];s.font.name='Times New Roman';s.font.size=Pt(size)
    s.font.bold=bold;s.font.color.rgb=RGBColor(0,0,0)
    p=s.paragraph_format;p.line_spacing=1.5;p.space_after=Pt(6);p.space_before=Pt(0)
    return s
body=style('Results Body');body.paragraph_format.first_line_indent=Inches(.25)
heading=style('Results Heading',12,True);heading.paragraph_format.keep_with_next=True
heading.paragraph_format.space_before=Pt(10)
table_style=style('Results Table',10.5);table_style.paragraph_format.line_spacing=1.1
table_style.paragraph_format.space_after=Pt(0)
caption_style=style('Results Caption',11,True);caption_style.paragraph_format.keep_with_next=True
caption_style.paragraph_format.line_spacing=1.15
def para(text):
    p=doc.add_paragraph(text,'Results Body');p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    return p
def h(text):return doc.add_paragraph(text,'Results Heading')
def table(caption,headers,rows,widths):
    doc.add_paragraph(caption,'Results Caption')
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for cell,width in zip(t.columns,widths):cell.width=Inches(width)
    for i,text in enumerate(headers):t.rows[0].cells[i].text=text
    for row in rows:
        for cell,text in zip(t.add_row().cells,row):cell.text=str(text)
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    for ri,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if ri==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'404040' if ri==0 else ('F3F3F3' if ri%2==0 else 'FFFFFF'));pr.append(shade)
            margins=OxmlElement('w:tcMar')
            for edge in ('top','left','bottom','right'):
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
            for p in cell.paragraphs:
                p.style=table_style;p.paragraph_format.keep_with_next=False
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.bold=ri==0;r.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
    doc.add_paragraph('', 'Results Table').paragraph_format.space_after=Pt(4)

title=doc.add_paragraph('Chapter 4 Results and Discussion','Title')
title.paragraph_format.left_indent=Inches(0);title.paragraph_format.first_line_indent=Inches(0)
para('The generator matched the C reference exactly across the automated comparison workload. All four Assembly waveforms met the generation-time target on the test laptop, although optimized C remained faster. This chapter presents the implementation results and the limits of the evaluation.')
h('4.1 Implemented System and Evaluation Conditions')
para('The program combines a hand-written x86-64 NASM engine with a Python and Tkinter interface. The engine exports sample counting, sample generation, and WAV construction through sfx.dll. A ctypes binding allocates the buffers and passes the parameter structure to the engine. The frontend provides four waveform choices, nine parameter sliders, a noise-seed field, seven presets, Randomize, a waveform view with a 20 ms zoom option, Play, Save WAV, and an auto-play toggle that starts disabled.')
para('Windows x64, Tkinter, threaded winsound playback, and MinGW-w64 were selected for this implementation. A single build script produces the engine, C reference libraries, native test harness, and benchmark executable. Testing was performed on one laptop; the results therefore do not establish performance across the whole group\'s machines.')
table('Table 4.1 Evaluation environment',['Item','Recorded configuration'],[
('Processor',cpu['Name'].strip()),('CPU resources',f"{cpu['NumberOfCores']} cores and {cpu['NumberOfLogicalProcessors']} logical processors"),
('Operating system','Windows 11 x64 build 26200'),('Assembly and C tools','NASM 2.16.03 and GCC 15.2.0'),
('Python','3.12.14 for binding tests; 3.11.9 for Tkinter tests'),('Timing power scheme','Balanced'),
('Timing conditions','Background apps and AC connection not controlled')],[1.55,4.45])
para('Measurements were recorded on 26 October 2026. The C reference used independently written synthesis routines, identical binary32 sine constants, and disabled floating-point contraction to preserve rounding. Agreement with this reference verifies the specified algorithm but does not establish perceptual quality or exclude a shared interpretation error.')
h('4.2 Correctness and Interface Results')
para(f"The native harness completed {native['checks']:,} assertions with no failures. It compared 20 fixed parameter sets, five for each waveform and including all seven presets, plus 10,000 reproducibly seeded random valid sets. Every comparison returned the same sample count and identical signed 16-bit samples. The maximum observed sample difference was {native['maximum_sample_delta']} LSB, meeting the allowed difference of one LSB.")
table('Table 4.2 Automated verification results',['Check group','Workload or condition','Observed result'],[
('Reference comparison','20 fixed and 10,000 random sets','Identical samples'),
('Repeatability','1,000 repeated native calls','Identical bytes'),
('Input validation','Range, NaN, infinity, null and cross-field errors','Rejected without writes'),
('Buffer protection','Short, exact and oversized capacities','Guard bytes intact'),
('ABI preservation','Success and error paths','Nonvolatile registers preserved'),
('WAV construction','1, 44,100 and 220,500 samples','Header and payload matched'),
('Randomize','1,000 binding-level runs','Valid sounds generated'),
('Concurrent generation','100 calls using four worker threads','Identical expected output')],[1.5,2.6,1.9])
para('Native checks also covered pitch-slide timing, early termination, waveform periods, full-scale quantization, and maximum upward slide. A descending sound whose minimum frequency equals its starting frequency generated one sample. WAV construction uses the actual generated count, preventing unused sample-buffer contents from appearing after early termination.')
para('PE inspection found exactly three exports, no imported DLLs, and no writable sections. The read-only import directory contains only its null terminator; Windows unwind metadata also remains read-only. The 44-byte ctypes parameter layout and every field offset matched the C and Assembly interface.')
h('4.3 Generated Sound Files')
para('All seven preset WAV files were read back with Python\'s wave module. Each reported mono audio at 44,100 Hz with a two-byte sample width, and its payload matched the generated samples exactly. Table 4.3 reports the outputs for seed 1.')
table('Table 4.3 Exported preset outputs',['Preset','Samples','Duration in s','WAV bytes'],[
    (r['preset'],f"{r['samples']:,}",f"{r['duration_s']:.2f}",f"{r['wav_bytes']:,}") for r in headless['presets']],[1.8,1.4,1.4,1.4])
para('Repeated generation reproduced each preset file exactly, while changing the noise seed changed the noise output. These results establish structural WAV correctness and repeatability. Preset recognisability requires the group listening test; Audacity and Windows Media Player compatibility also remain unverified.')
h('4.4 Generation Performance and Resource Use')
para('The native benchmark measured synthesis alone using QueryPerformanceCounter. Each workload generated one second of audio, or 44,100 samples, with pitch slide and vibrato enabled. For each waveform and implementation, 100 warm-up calls preceded 1,000 measured calls. The output buffer was allocated before timing. DLL loading, WAV construction, Python conversion, canvas drawing, playback, and file writing were excluded. Table 4.4 reports medians for all implementations and the Assembly 95th percentile.')
table('Table 4.4 Time to generate 44,100 samples in milliseconds',
['Waveform','Assembly median','Assembly p95','C O2 median','C O0 median'],[
    (waves[w],f"{lookup['Assembly',w]['median_ms']:.3f}",f"{lookup['Assembly',w]['p95_ms']:.3f}",f"{lookup['C O2',w]['median_ms']:.3f}",f"{lookup['C O0',w]['median_ms']:.3f}") for w in range(4)
],[1.2,1.2,1.2,1.2,1.2])
minimum=min(lookup['Assembly',w]['median_ms'] for w in range(4));maximum=max(lookup['Assembly',w]['median_ms'] for w in range(4))
ratios=[lookup['Assembly',w]['median_ms']/lookup['C O2',w]['median_ms'] for w in range(4)]
para(f'Assembly medians ranged from {minimum:.3f} to {maximum:.3f} ms, below the 5 ms target for every waveform on this machine. However, Assembly took {min(ratios):.2f} to {max(ratios):.2f} times the optimized C median. It was faster than unoptimized C in all four measured workloads. The measurements therefore support the speed requirement on this laptop but do not support a claim that hand-written Assembly is inherently faster than optimized C.')
para('Branch structure, scalar dependencies, and instruction choices may explain the timing differences, but this remains a hypothesis without hardware-counter analysis or a controlled optimization experiment. Fixed execution order and uncontrolled background activity also limit interpretation of small differences.')
para(f"The stripped Assembly DLL occupied {headless['dll_bytes']:,} bytes, below the 20 KB target. The optimized C reference object occupied {env['c_reference_object_bytes']:,} bytes; an object and a complete DLL are different packaging units, so these sizes do not establish a language-level size advantage. Source inspection found a 152-byte generation stack frame, no heap allocation, and no writes to global state. Including nested return addresses, engine stack use remains below 256 bytes. The NASM listing contains {loop_instructions} statically emitted instructions between the loop and completion labels across alternative paths; this is not the dynamic instruction count for each sample.")
h('4.5 Frontend Integration Results')
para(f"The binding suite passed {headless['checks']:,} assertions. A separate off-screen Tkinter suite passed {gui['checks']} assertions using real widgets and a 480-column Canvas. It exercised every preset, slider limits, seed validation, zoom, saving the current WAV, and the missing-DLL error state. Playback was mocked to verify that the callback runs in a worker thread and the UI event loop continues processing events. This test did not send sound to the speakers.")
para(f"For 100 maximum-length sounds of 220,500 samples, the off-screen test measured a median event-to-canvas time of {gui['median_event_to_canvas_ms']:.1f} ms and a 95th percentile of {gui['p95_event_to_canvas_ms']:.1f} ms, including the 50 ms debounce. The measured p95 was below the 100 ms target in this setup. The withdrawn window and fixed Canvas dimensions exclude visible-window composition and display costs; a visible interactive measurement is still required before treating NFR-04 as fully verified.")
h('4.6 Remaining Evaluation and Interpretation')
para('Several acceptance activities remain open. A second group member must review the visible interface, operate Play and Save, and confirm speaker output. The group must listen to all presets and judge whether each effect is recognisable. WAV files still need to be opened in Audacity and Windows Media Player. The ten-minute interactive soak session, clean-install check, and benchmark repetition on a second laptop have not been completed. The 1,000 automated Randomize calls provide stability evidence for generation, but do not replace those activities.')
para('Final benchmarking should use connected power, a recorded power mode, closed background applications, and a second laptop. A later battery snapshot cannot establish power conditions throughout timing. Raw test output, environment metadata, disassembly, and source and DLL hashes are retained in the results directory for reproduction.')
para('The system demonstrates reproducible Assembly synthesis, frontend integration, and PCM WAV export within the measured resource targets on one machine. Independent user review, listening tests, external-player checks, longer stability testing, and performance replication are required to complete the evaluation.')
doc.save(target)
print(target)
print('Original preserved at',backup)
print('Loop static instructions',loop_instructions)
